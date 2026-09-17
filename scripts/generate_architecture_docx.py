import os
import re
import subprocess
from PIL import Image
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
docs_dir = os.path.join(base_dir, "docs")
html_path = os.path.join(docs_dir, "Spesifikasi_Teknis_Arsitektur_dan_Alur_Logika_Helpdesk_GTT.html")
docx_path = os.path.join(docs_dir, "Rancangan_Arsitektur_dan_Alur_Helpdesk_Ticketing.docx")
edge_path = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
if not os.path.exists(edge_path):
    edge_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

# --- 1. RENDER SVG DIAGRAMS TO HIGH-RES PNG IMAGES ---
with open(html_path, "r", encoding="utf-8") as f:
    html_content = f.read()

svgs = re.findall(r"(<svg.*?</svg>)", html_content, re.DOTALL)
print(f"Total SVGs found in HTML: {len(svgs)}")

arch_png = os.path.join(docs_dir, "diagram_arsitektur.png")
flow_png = os.path.join(docs_dir, "diagram_flowchart.png")

for i, svg in enumerate(svgs[:2]):
    tag = "arsitektur" if i == 0 else "flowchart"
    target_png = arch_png if i == 0 else flow_png
    temp_html = os.path.join(docs_dir, f"temp_{tag}.html")
    temp_screenshot = os.path.join(docs_dir, f"temp_{tag}.png")

    page_markup = f"""<!DOCTYPE html>
<html>
<head>
  <meta charset="UTF-8">
  <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;600;700&display=swap');
    body {{
      margin: 0;
      padding: 10px;
      background: #ffffff;
      font-family: 'Inter', sans-serif;
      display: inline-block;
    }}
    svg {{
      width: 1000px;
      height: auto;
      display: block;
    }}
  </style>
</head>
<body>
{svg}
</body>
</html>"""
    
    with open(temp_html, "w", encoding="utf-8") as f_tmp:
        f_tmp.write(page_markup)

    cmd = [
        edge_path,
        "--headless=new",
        "--disable-gpu",
        "--window-size=1100,500",
        f"--screenshot={temp_screenshot}",
        temp_html
    ]
    subprocess.run(cmd, capture_output=True, text=True)

    if os.path.exists(temp_screenshot):
        im = Image.open(temp_screenshot)
        bbox = im.getbbox()
        if bbox:
            cropped = im.crop(bbox)
            cropped.save(target_png)
            print(f"Rendered {target_png} ({cropped.size[0]}x{cropped.size[1]} px)")
        else:
            im.save(target_png)
        os.remove(temp_screenshot)
    if os.path.exists(temp_html):
        os.remove(temp_html)

# --- 2. BUILD THE WORD DOCUMENT (.DOCX) ---
doc = docx.Document()

# Page Margins (Standard A4)
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

# Default style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)

# --- HEADER / TITLE ---
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(2)
r_t = p_title.add_run("RANCANGAN ARSITEKTUR, MATRIKS AKSES, DAN ALUR SISTEM HELPDESK TICKETING")
r_t.bold = True
r_t.font.size = Pt(13.5)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(12)
r_s = p_sub.add_run("Dokumen Spesifikasi Teknis: Desain Arsitektur Full-Stack (Vue.js, Express.js, MySQL, SMTP), Matriks Hak Akses (RBAC), Flowchart Penanganan Insiden, dan Logika Evaluasi SLA")
r_s.font.size = Pt(9.5)
r_s.italic = True
r_s.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

# --- DAFTAR SINGKATAN & GLOSARIUM ---
p_glo_h = doc.add_paragraph()
p_glo_h.paragraph_format.space_before = Pt(6)
p_glo_h.paragraph_format.space_after = Pt(4)
r_gh = p_glo_h.add_run("DAFTAR ISTILAH DAN SINGKATAN")
r_gh.bold = True
r_gh.font.size = Pt(10.5)

t_glo = doc.add_table(rows=1, cols=3)
t_glo.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_glo)

hdr_g = t_glo.rows[0].cells
hdr_g[0].text = "Singkatan"
hdr_g[1].text = "Kepanjangan / Istilah Lengkap"
hdr_g[2].text = "Penjelasan Fungsi dalam Sistem"
for c in hdr_g:
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)

glo_rows = [
    ("CPIG", "Customer Problem Incident Group", "Staf Helpdesk & Dispatcher yang menerima keluhan, menerbitkan tiket, mengalokasikan teknisi, dan berkomunikasi dengan klien."),
    ("SLA", "Service Level Agreement", "Kesepakatan batas waktu layanan penanganan insiden yang mengikat secara kontraktual."),
    ("MTTR", "Mean Time To Resolve", "Rata-rata durasi waktu yang dibutuhkan dari insiden tercatat hingga sistem normal kembali."),
    ("Net MTTR", "Net Mean Time To Resolve", "Durasi penyelesaian bersih setelah dikurangi seluruh akumulasi jeda waktu resmi (SLA Pause)."),
    ("RBAC", "Role-Based Access Control", "Metode pembatasan hak akses sistem berdasarkan peran kerja pengguna."),
    ("RCA", "Root Cause Analysis", "Investigasi penelusuran akar penyebab utama terjadinya gangguan teknis."),
    ("KCS", "Knowledge-Centered Service", "Metodologi standardisasi perbaikan insiden menjadi dokumen panduan teknis (SOP / Runbook)."),
    ("SMTP", "Simple Mail Transfer Protocol", "Protokol standar jaringan untuk pengiriman surat elektronik resmi."),
    ("SPA", "Single Page Application", "Arsitektur web yang memuat satu halaman dinamis tanpa memuat ulang seluruh halaman peramban."),
    ("PKS", "Perjanjian Kerja Sama", "Dokumen kontrak tingkat layanan antara penyedia layanan dengan mitra perusahaan.")
]

for idx, item in enumerate(glo_rows):
    row = t_glo.add_row()
    for c_idx, val in enumerate(item):
        cell = row.cells[c_idx]
        cell.text = val
        if idx % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# --- 1. ARSITEKTUR SISTEM FULL-STACK ---
p_s1 = doc.add_paragraph()
p_s1.paragraph_format.space_before = Pt(10)
p_s1.paragraph_format.space_after = Pt(4)
r_s1 = p_s1.add_run("1. ARSITEKTUR SISTEM FULL-STACK")
r_s1.bold = True
r_s1.font.size = Pt(11.5)

p_s1_desc = doc.add_paragraph()
p_s1_desc.paragraph_format.space_after = Pt(6)
p_s1_desc.add_run(
    "Sistem dirancang dengan arsitektur berlapis (multi-tier architecture) yang menghubungkan antarmuka pengguna berbasis "
    "Vue.js 3, backend pemrosesan logika bisnis berbasis Express.js (Node.js), basis data relasional MySQL, dan "
    "pengiriman notifikasi surat elektronik melalui SMTP Relay:"
)

# Insert Picture of Architecture Diagram
if os.path.exists(arch_png):
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_before = Pt(4)
    p_img1.paragraph_format.space_after = Pt(4)
    p_img1.add_run().add_picture(arch_png, width=Inches(6.4))
    p_cap1 = doc.add_paragraph()
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap1.paragraph_format.space_after = Pt(8)
    r_cap1 = p_cap1.add_run("Gambar 1.1 Diagram Arsitektur Sistem Full-Stack (Vue.js, Express.js, MySQL, SMTP Relay)")
    r_cap1.font.size = Pt(8.5)
    r_cap1.italic = True
    r_cap1.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

t_arch = doc.add_table(rows=1, cols=3)
t_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_arch)

hdr_a = t_arch.rows[0].cells
hdr_a[0].text = "Komponen Sistem"
hdr_a[1].text = "Teknologi Terpilih"
hdr_a[2].text = "Peran & Fungsi dalam Rancangan"
for c in hdr_a:
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)

arch_rows = [
    ("Lapisan Antarmuka", "Vue 3, Vite, Pinia, Vue Router", "Menyediakan antarmuka interaktif bagi staf Helpdesk, lembar kerja diagnosa teknisi, dan portal tracking publik bagi pelanggan."),
    ("Layanan Backend", "Express.js, Node-Cron, Nodemailer", "Memproses endpoint REST API, menjalankan validasi RBAC, menghitung batas SLA, serta mengeksekusi cron monitoring berkala."),
    ("Basis Data Relasional", "MySQL (Engine InnoDB)", "Menyimpan seluruh tabel operasional tiket, relasi pelanggan, kebijakan SLA, riwayat jeda, dan catatan rekam audit."),
    ("Layanan Notifikasi", "SMTP Relay Server", "Mengirimkan pembaruan status tiket dan surat peringatan eskalasi resmi langsung ke kotak masuk surel teknisi dan klien.")
]

for idx, item in enumerate(arch_rows):
    row = t_arch.add_row()
    for c_idx, val in enumerate(item):
        cell = row.cells[c_idx]
        cell.text = val
        if idx % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True

doc.add_page_break()

# --- 2. MATRIKS HAK AKSES PENGGUNA (RBAC) ---
p_s2 = doc.add_paragraph()
p_s2.paragraph_format.space_before = Pt(8)
p_s2.paragraph_format.space_after = Pt(4)
r_s2 = p_s2.add_run("2. MATRIKS HAK AKSES PENGGUNA (RBAC MATRIX)")
r_s2.bold = True
r_s2.font.size = Pt(11.5)

p_s2_desc = doc.add_paragraph()
p_s2_desc.paragraph_format.space_after = Pt(6)
p_s2_desc.add_run(
    "Pembagian hak akses memisahkan tugas ADMIN / CPIG (sebagai Helpdesk Dispatcher yang mengatur antrean dan penugasan) "
    "dengan SUPPORT ENGINEER (sebagai teknisi pelaksana investigasi dan perbaikan), serta membatasi akses PORTAL KLIEN:"
)

t_rbac = doc.add_table(rows=1, cols=4)
t_rbac.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_rbac)

hdr_r = t_rbac.rows[0].cells
hdr_r[0].text = "Fungsi / Modul Sistem"
hdr_r[1].text = "ADMIN / CPIG\n(Helpdesk Dispatcher)"
hdr_r[2].text = "SUPPORT ENGINEER\n(Teknisi Penanganan)"
hdr_r[3].text = "PORTAL KLIEN\n(Pelacakan Mandiri)"
for c in hdr_r:
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    for p in c.paragraphs:
        if c != hdr_r[0]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)

rbac_rows = [
    ("Monitoring Dashboard & Antrean Global Tiket", "Akses Penuh", "Tiket Ditugaskan", "Ditolak", True, False, False),
    ("Penerbitan Tiket Baru & Pengisian Data Awal", "Akses Penuh", "Input Lapangan", "Ditolak", True, False, False),
    ("Penugasan & Pergantian Teknisi Penanggung Jawab", "Kontrol Penuh", "Hanya Baca", "Ditolak", True, False, False),
    ("Penyalinan Tautan Pelacakan Klien (Token URL)", "Akses Khusus", "Disembunyikan", "Ditolak", True, False, False),
    ("Pratinjau & Pengiriman Surel Notifikasi Manual", "Akses Khusus", "Disembunyikan", "Ditolak", True, False, False),
    ("Pencatatan Milestone & Log Perintah Perbaikan", "Bisa Menambah", "Pelaksana Utama", "Ditolak", False, True, False),
    ("Pengajuan & Pencabutan Jeda SLA (Pause)", "Validasi Otoritas", "Pelaksana Utama", "Ditolak", False, True, False),
    ("Penyelesaian Tiket & Analisis Akar Masalah (RCA)", "Verifikasi Akhir", "Pelaksana Utama", "Ditolak", False, True, False),
    ("Pengaturan Matriks SLA & Data Klien", "Akses Penuh", "Ditolak", "Ditolak", True, False, False),
    ("Laporan SLA & Ekspor Berkas Laporan", "Akses Penuh", "Ditolak", "Ditolak", True, False, False),
    ("Basis Pengetahuan (Knowledge Base)", "Persetujuan SOP", "Ajukan & Membaca", "Hanya Baca", True, False, False)
]

for idx, (mod, adm, tch, clt, b_adm, b_tch, b_clt) in enumerate(rbac_rows):
    row = t_rbac.add_row()
    row.cells[0].text = mod
    row.cells[1].text = adm
    row.cells[2].text = tch
    row.cells[3].text = clt
    
    if idx % 2 == 1:
        for c in row.cells:
            set_cell_shading(c, "F8FAFC")
            
    for c_idx, cell in enumerate(row.cells):
        set_cell_margins(cell, top=55, bottom=55, left=90, right=90)
        for p in cell.paragraphs:
            if c_idx > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(8.5)
                if c_idx == 1 and b_adm:
                    r.bold = True
                elif c_idx == 2 and b_tch:
                    r.bold = True
                elif c_idx == 3 and b_clt:
                    r.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# --- 3. DIAGRAM ALUR PROSES (FLOWCHART) ---
p_s3 = doc.add_paragraph()
p_s3.paragraph_format.space_before = Pt(10)
p_s3.paragraph_format.space_after = Pt(4)
r_s3 = p_s3.add_run("3. DIAGRAM ALUR PROSES PENANGANAN TIKET (FLOWCHART)")
r_s3.bold = True
r_s3.font.size = Pt(11.5)

p_s3_desc = doc.add_paragraph()
p_s3_desc.paragraph_format.space_after = Pt(6)
p_s3_desc.add_run(
    "Siklus hidup setiap insiden dimodelkan dalam bentuk diagram alur proses terstruktur dari pelaporan awal hingga tiket ditutup:"
)

# Insert Picture of Flowchart Diagram
if os.path.exists(flow_png):
    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_before = Pt(4)
    p_img2.paragraph_format.space_after = Pt(4)
    p_img2.add_run().add_picture(flow_png, width=Inches(6.4))
    p_cap2 = doc.add_paragraph()
    p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap2.paragraph_format.space_after = Pt(8)
    r_cap2 = p_cap2.add_run("Gambar 1.2 Diagram Alur Logika Penanganan Tiket (Incident Lifecycle Flowchart)")
    r_cap2.font.size = Pt(8.5)
    r_cap2.italic = True
    r_cap2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

flow_steps = [
    ("1. Laporan Masuk & Pencatatan Tiket", "Petugas Helpdesk/CPIG mencatat tiket baru ke sistem Express.js. Sistem mengunci target respon dan resolusi SLA secara otomatis dari tabel master SLA sesuai kontrak pelanggan. Email konfirmasi dan tautan token dibuat dan dikirimkan via SMTP."),
    ("2. Respon Awal Teknisi (First Touch SLA)", "Teknisi mengubah status tiket menjadi IN_PROGRESS. Sistem menghentikan timer perhitungan Response SLA dan mencatat waktu respon pertama secara permanen di database."),
    ("3. Investigasi Diagnostik & Milestone", "Teknisi melakukan penelusuran masalah dan mencatatkan setiap tindakan teknis, hasil observasi, serta riwayat eksekusi perintah terminal ke dalam tabel ticket_milestones."),
    ("4. Penanganan Jeda Waktu SLA (Clock Pause)", "Jika perbaikan terhenti akibat menunggu komponen suku cadang dari vendor principal atau menunggu jendela pemeliharaan klien, teknisi mengaktifkan status Pause. Timer resolusi SLA dibekukan sementara sehingga sisa waktu SLA tidak terpotong."),
    ("5. Pemulihan Sistem & Pengisian RCA", "Setelah infrastruktur kembali normal, teknisi mengisi formulir penyelesaian yang memuat penyebab gangguan (Root Cause), langkah pemulihan, dan saran preventif. Tiket diubah ke status RESOLVED dan waktu resolusi bersih (Net MTTR) dikunci."),
    ("6. Verifikasi Klien & Penutupan Tiket (Closed)", "Prosedur perbaikan yang efektif dapat diajukan ke antrean persetujuan Knowledge Base (KCS). Tiket ditutup permanen (CLOSED) setelah adanya konfirmasi dari pelanggan atau melewati masa evaluasi.")
]

for title, desc in flow_steps:
    p_st = doc.add_paragraph()
    p_st.paragraph_format.space_before = Pt(2)
    p_st.paragraph_format.space_after = Pt(3)
    p_st.paragraph_format.left_indent = Inches(0.2)
    rt = p_st.add_run(f"• {title}: ")
    rt.bold = True
    rt.font.size = Pt(9)
    rd = p_st.add_run(desc)
    rd.font.size = Pt(9)

doc.add_page_break()

# --- 4. FORMULA & MATRIKS SLA ---
p_s4 = doc.add_paragraph()
p_s4.paragraph_format.space_before = Pt(8)
p_s4.paragraph_format.space_after = Pt(4)
r_s4 = p_s4.add_run("4. FORMULA & MATRIKS PERHITUNGAN SLA (SERVICE LEVEL AGREEMENT)")
r_s4.bold = True
r_s4.font.size = Pt(11.5)

p_s4_desc = doc.add_paragraph()
p_s4_desc.paragraph_format.space_after = Pt(6)
p_s4_desc.add_run(
    "Evaluasi performa layanan memisahkan dua parameter utama secara independen pada backend Express.js:"
)

# Formula Cards Table
t_form = doc.add_table(rows=1, cols=2)
t_form.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_form)

fc1 = t_form.rows[0].cells[0]
fc2 = t_form.rows[0].cells[1]
set_cell_shading(fc1, "F8FAFC")
set_cell_shading(fc2, "F8FAFC")
set_cell_margins(fc1, top=80, bottom=80, left=120, right=120)
set_cell_margins(fc2, top=80, bottom=80, left=120, right=120)

p1 = fc1.paragraphs[0]
r = p1.add_run("A. Formula Waktu Respon Awal (First Touch SLA)\n")
r.bold = True
r.font.size = Pt(9)
r2 = p1.add_run("ΔT_respon = T_respon - T_dibuat\n")
r2.bold = True
r2.font.name = 'Consolas'
r2.font.size = Pt(9.5)
r3 = p1.add_run("Kriteria Terpenuhi: ΔT_respon ≤ Batas Respon Kontrak PKS")
r3.font.size = Pt(8)
r3.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

p2 = fc2.paragraphs[0]
r = p2.add_run("B. Formula Waktu Resolusi Bersih (Net MTTR SLA)\n")
r.bold = True
r.font.size = Pt(9)
r2 = p2.add_run("Net MTTR = (T_selesai - T_dibuat) - Σ T_pause\n")
r2.bold = True
r2.font.name = 'Consolas'
r2.font.size = Pt(9.5)
r3 = p2.add_run("Kriteria Terpenuhi: Net MTTR ≤ Batas Resolusi Kontrak PKS")
r3.font.size = Pt(8)
r3.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Matriks PKS
p_pks_h = doc.add_paragraph()
p_pks_h.paragraph_format.space_before = Pt(6)
p_pks_h.paragraph_format.space_after = Pt(4)
r_pks_h = p_pks_h.add_run("A. Matriks Standar Kontrak Tingkat Layanan (PKS)")
r_pks_h.bold = True
r_pks_h.font.size = Pt(10)

t_pks = doc.add_table(rows=1, cols=4)
t_pks.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_pks)

hdr_p = t_pks.rows[0].cells
hdr_p[0].text = "Tingkatan Kontrak"
hdr_p[1].text = "Jadwal Jam Operasional"
hdr_p[2].text = "Target Waktu Respon"
hdr_p[3].text = "Target Resolusi Net MTTR"
for c in hdr_p:
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    for p in c.paragraphs:
        if c != hdr_p[0]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)

pks_rows = [
    ("Platinum 24×7", "24 Jam × 7 Hari (Non-Stop)", "≤ 30 Menit", "≤ 4 Jam"),
    ("Gold 8×5", "Senin – Jumat (08:00 – 17:00 WIB)", "≤ 30 Menit", "≤ 8 Jam"),
    ("Silver 8×5", "Senin – Jumat (08:30 – 17:30 WIB)", "≤ 1 Jam", "≤ 12 Jam"),
    ("Pemerintah Tier-1", "Jam Kantor Dinas Pemerintah (8×5)", "≤ 1 Jam", "≤ 12 Jam")
]

for idx, item in enumerate(pks_rows):
    row = t_pks.add_row()
    for c_idx, val in enumerate(item):
        cell = row.cells[c_idx]
        cell.text = val
        if idx % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, top=55, bottom=55, left=90, right=90)
        for p in cell.paragraphs:
            if c_idx >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Matriks Eskalasi
p_esk_h = doc.add_paragraph()
p_esk_h.paragraph_format.space_before = Pt(6)
p_esk_h.paragraph_format.space_after = Pt(4)
r_esk_h = p_esk_h.add_run("B. Matriks Ambang Batas Pemicu Eskalasi Otomatis (Node-Cron Service)")
r_esk_h.bold = True
r_esk_h.font.size = Pt(10)

t_esk = doc.add_table(rows=1, cols=4)
t_esk.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_esk)

hdr_e = t_esk.rows[0].cells
hdr_e[0].text = "Tahap Eskalasi"
hdr_e[1].text = "Ambang Batas Waktu"
hdr_e[2].text = "Aksi Sistem Otomatis"
hdr_e[3].text = "Tujuan Notifikasi Surel"
for c in hdr_e:
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)

esk_rows = [
    ("1. Peringatan Dini", "Durasi Mencapai 50% Target MTTR", "Kirim surel pengingat status pengerjaan", "Teknisi Penanggung Jawab & Lead CPIG"),
    ("2. Peringatan Kritis", "Durasi Mencapai 80% Target MTTR", "Kirim surel alert darurat (High Priority)", "Service Delivery Lead / Supervisor"),
    ("3. Pelanggaran SLA", "100% Target MTTR Terlampaui", "Ubah status menjadi Breached & catat audit log", "Manajemen Operasional Helpdesk")
]

for idx, item in enumerate(esk_rows):
    row = t_esk.add_row()
    for c_idx, val in enumerate(item):
        cell = row.cells[c_idx]
        cell.text = val
        if idx % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, top=55, bottom=55, left=90, right=90)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Skema Database
p_db_h = doc.add_paragraph()
p_db_h.paragraph_format.space_before = Pt(6)
p_db_h.paragraph_format.space_after = Pt(4)
r_db_h = p_db_h.add_run("C. Rincian Entitas Skema Basis Data (MySQL)")
r_db_h.bold = True
r_db_h.font.size = Pt(10)

t_db = doc.add_table(rows=1, cols=3)
t_db.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_db)

hdr_d = t_db.rows[0].cells
hdr_d[0].text = "Tabel Basis Data"
hdr_d[1].text = "Kolom Utama"
hdr_d[2].text = "Fungsi Penyimpanan"
for c in hdr_d:
    set_cell_shading(c, "F1F5F9")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)

db_rows = [
    ("tickets", "id, ticket_number, customer_id, assigned_to, status, priority, sla_policy_id", "Menyimpan data tiket insiden dan relasi ke pelanggan serta teknisi penanggung jawab."),
    ("sla_pause_logs", "id, ticket_id, reason_type, vendor_case_id, paused_at, resumed_at, total_seconds", "Mencatat histori jeda waktu resmi (menunggu suku cadang vendor atau izin jadwal klien)."),
    ("ticket_milestones", "id, ticket_id, title, description, command_executed, created_by, created_at", "Rekam jejak kronologis langkah diagnosa teknis dan eksekusi perintah terminal oleh teknisi."),
    ("knowledge_base", "id, kb_number, title, category, target_vendor, runbook_steps, status, approved_by", "Katalog dokumen SOP terverifikasi yang dihasilkan dari daur ulang solusi insiden nyata."),
    ("audit_event_logs", "id, ticket_id, user_id, action_type, old_value, new_value, timestamp", "Catatan jejak audit permanen untuk menjamin integritas data dan riwayat penanganan.")
]

for idx, item in enumerate(db_rows):
    row = t_db.add_row()
    for c_idx, val in enumerate(item):
        cell = row.cells[c_idx]
        cell.text = val
        if idx % 2 == 1:
            set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, top=55, bottom=55, left=90, right=90)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True
                    r.font.name = 'Consolas'

# Save file
doc.save(docx_path)
print(f"Publication-grade Word document generated successfully at: {docx_path}")
